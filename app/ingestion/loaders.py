from __future__ import annotations
from pathlib import Path
import json, hashlib, re
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image
import pytesseract
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from app.models import Source, Chunk

TRUST = {
    'manuals': ('official_manual', 1.00),
    'engineering_bulletins': ('engineering_change', 0.98),
    'reference': ('official_reference', 0.95),
    'configuration': ('configuration_export', 0.93),
    'diagrams': ('engineering_diagram', 0.90),
    'screenshots': ('hmi_evidence', 0.85),
    'scans': ('scanned_record', 0.75),
    'extra': ('training_material', 0.70),
    'low_trust': ('field_notes', 0.35),
    'noise': ('irrelevant_reference', 0.10),
}

def source_info(path: Path) -> tuple[str,float]:
    for parent in path.parts[::-1]:
        if parent in TRUST:
            return TRUST[parent]
    return ('unknown', 0.5)

def sid(path: Path) -> str:
    return 'src_' + hashlib.sha1(str(path).encode()).hexdigest()[:10]

def cid(source_id, locator, text):
    return 'chk_' + hashlib.sha1(f'{source_id}|{locator}|{text[:100]}'.encode()).hexdigest()[:12]

def clean(text):
    return re.sub(r'\s+', ' ', text).strip()

def pdf_chunks(path: Path, source: Source):
    out=[]
    doc=fitz.open(path)
    for i,page in enumerate(doc):
        txt=page.get_text('text')
        if txt.strip():
            out.append(Chunk(cid(source.source_id,f'page:{i+1}',txt),source.source_id,f'page:{i+1}',clean(txt)))
        else:
            pix=page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
            img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
            ocr=pytesseract.image_to_string(img)
            if ocr.strip():
                out.append(Chunk(cid(source.source_id,f'page:{i+1}:ocr',ocr),source.source_id,f'page:{i+1}:ocr',clean(ocr),'ocr'))
    return out

def html_chunks(path,source):
    soup=BeautifulSoup(path.read_text(encoding='utf-8',errors='ignore'),'html.parser')
    parts=[]
    for tag in soup.find_all(['h1','h2','h3','p','li','tr']):
        t=clean(tag.get_text(' ',strip=True))
        if t: parts.append(t)
    return [Chunk(cid(source.source_id,f'html:{i+1}',t),source.source_id,f'html:{i+1}',t) for i,t in enumerate(parts)]

def docx_chunks(path,source):
    d=Document(path); out=[]
    for i,p in enumerate(d.paragraphs):
        t=clean(p.text)
        if t: out.append(Chunk(cid(source.source_id,f'paragraph:{i+1}',t),source.source_id,f'paragraph:{i+1}',t))
    for ti,tbl in enumerate(d.tables,1):
        rows=[]
        for r in tbl.rows: rows.append(' | '.join(clean(c.text) for c in r.cells))
        text='\n'.join(rows)
        if text: out.append(Chunk(cid(source.source_id,f'table:{ti}',text),source.source_id,f'table:{ti}',text,'table'))
    return out

def xlsx_chunks(path,source):
    wb=load_workbook(path,data_only=True); out=[]
    for ws in wb.worksheets:
        for r,row in enumerate(ws.iter_rows(values_only=True),1):
            vals=[str(v) for v in row if v is not None]
            if vals:
                t=' | '.join(vals)
                out.append(Chunk(cid(source.source_id,f'sheet:{ws.title}:row:{r}',t),source.source_id,f'sheet:{ws.title}:row:{r}',t,'table',{'sheet':ws.title,'row':r}))
    return out

def json_chunks(path,source):
    data=json.loads(path.read_text(encoding='utf-8'))
    out=[]
    def walk(obj,prefix=''):
        if isinstance(obj,dict):
            for k,v in obj.items(): walk(v,f'{prefix}.{k}' if prefix else k)
        elif isinstance(obj,list):
            for i,v in enumerate(obj): walk(v,f'{prefix}[{i}]')
        else:
            t=f'{prefix} = {obj}'
            out.append(Chunk(cid(source.source_id,prefix,t),source.source_id,prefix,t,'structured'))
    walk(data)
    return out

def pptx_chunks(path,source):
    prs=Presentation(path); out=[]
    for i,slide in enumerate(prs.slides,1):
        texts=[]
        for sh in slide.shapes:
            if hasattr(sh,'text') and sh.text.strip(): texts.append(clean(sh.text))
        t='\n'.join(texts)
        if t: out.append(Chunk(cid(source.source_id,f'slide:{i}',t),source.source_id,f'slide:{i}',t,'slide'))
    return out

def image_chunks(path,source):
    img=Image.open(path)
    txt=pytesseract.image_to_string(img)
    return [Chunk(cid(source.source_id,'image:ocr',txt),source.source_id,'image:ocr',clean(txt),'ocr')] if txt.strip() else []

def load_file(path: Path):
    trust,score=source_info(path)
    source=Source(sid(path),str(path),'/' .join(path.parts[-3:]),path.suffix.lower().lstrip('.'),trust,score)
    ext=path.suffix.lower()
    if ext=='.pdf': chunks=pdf_chunks(path,source)
    elif ext in {'.html','.htm'}: chunks=html_chunks(path,source)
    elif ext=='.docx': chunks=docx_chunks(path,source)
    elif ext=='.xlsx': chunks=xlsx_chunks(path,source)
    elif ext=='.json': chunks=json_chunks(path,source)
    elif ext=='.pptx': chunks=pptx_chunks(path,source)
    elif ext in {'.png','.jpg','.jpeg'}: chunks=image_chunks(path,source)
    else: chunks=[]
    return source,chunks

def ingest_tree(data_dir: str):
    root=Path(data_dir); sources=[]; chunks=[]
    for p in sorted(root.rglob('*')):
        if not p.is_file(): continue
        if p.suffix.lower() not in {'.pdf','.html','.htm','.docx','.xlsx','.json','.pptx','.png','.jpg','.jpeg'}: continue
        s,cs=load_file(p); sources.append(s); chunks.extend(cs)
    return sources,chunks
